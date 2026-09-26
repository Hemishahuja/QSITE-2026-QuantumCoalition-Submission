"""Generate a slide deck (PowerPoint) for the spoken-word demo in DEMO_SCRIPT.md.

Run from the `Computational Track` directory, after `generate_demo_assets.py`
has produced the PNGs this deck embeds:
    python solution/generate_demo_assets.py
    python solution/generate_slides.py

Writes solution/slides.pptx. Each slide has a short on-screen headline plus a
handful of bullets, and a "speaker notes" pane containing the exact spoken
sentences from DEMO_SCRIPT.md for that beat of the talk (word-for-word, so
the 594-word / 4:34 timing in that file still holds). Open the notes pane in
PowerPoint / Keynote / Google Slides ("Presenter View" / "Speaker Notes") and
read it straight through while screen-recording; nothing needs to be
memorized or improvised.

Bracketed lines inside notes (e.g. "[optional: ...]") are stage directions,
matching the convention already used in DEMO_SCRIPT.md -- do not read those
aloud.

This script does not call solve() or the scorer; it only lays out already
-verified numbers and already-generated images. If the numbers in
DEMO_SCRIPT.md change, update SLIDES below to match before regenerating.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "demo_assets"
OUT = HERE / "slides.pptx"

# -- palette, shared with the matplotlib assets (demo_assets/*.png) ---------
NAVY = RGBColor(0x1E, 0x29, 0x3B)       # table header / title bg
GREEN = RGBColor(0x16, 0xA3, 0x4A)      # "ours" / proven
AMBER = RGBColor(0xB4, 0x5F, 0x06)      # best-found / open (darker than the
                                        # chart's fill so it reads on white)
GRAY = RGBColor(0x47, 0x55, 0x69)       # baseline / secondary text
INK = RGBColor(0x1E, 0x29, 0x3B)        # body text
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF1, 0xF5, 0xF9)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.text_frame.word_wrap = True
    return box


def _kicker_and_title(slide, kicker: str, title: str, accent: RGBColor = NAVY) -> None:
    if kicker:
        kb = _textbox(slide, Inches(0.7), Inches(0.35), Inches(8), Inches(0.4))
        p = kb.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = kicker.upper()
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = accent
        run.font.name = "Calibri"

    tb = _textbox(slide, Inches(0.7), Inches(0.72), Inches(11.9), Inches(0.95))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.color.rgb = INK
    run.font.name = "Calibri"

    # thin accent rule under the title
    line = slide.shapes.add_shape(1, Inches(0.7), Inches(1.62), Inches(2.0), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = accent
    line.line.fill.background()
    line.shadow.inherit = False


def _bullets(slide, items, top=Inches(1.95), left=Inches(0.85), width=Inches(11.6),
             font_size=20, line_spacing=1.25, bold_prefix=False):
    box = _textbox(slide, left, top, width, Inches(7.3 - top.inches))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, tuple):
            text, color = item
        else:
            text, color = item, INK
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = line_spacing
        p.space_after = Pt(10)
        bullet_run = p.add_run()
        bullet_run.text = "\u25A0  "
        bullet_run.font.size = Pt(font_size)
        bullet_run.font.color.rgb = color if color != INK else GRAY
        bullet_run.font.bold = True
        bullet_run.font.name = "Calibri"
        text_run = p.add_run()
        text_run.text = text
        text_run.font.size = Pt(font_size)
        text_run.font.color.rgb = INK
        text_run.font.name = "Calibri"
    return box


def _image_fitted(slide, path: Path, top, left, max_w, max_h):
    from PIL import Image

    with Image.open(path) as im:
        px_w, px_h = im.size
    aspect = px_w / px_h
    w = max_w
    h = Emu(int(w / aspect))
    if h > max_h:
        h = max_h
        w = Emu(int(h * aspect))
    x = Emu(int(left + (max_w - w) / 2))
    y = Emu(int(top + (max_h - h) / 2))
    slide.shapes.add_picture(str(path), x, y, width=w, height=h)


def _notes(slide, text: str) -> None:
    tf = slide.notes_slide.notes_text_frame
    tf.text = text


def add_title_slide(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, NAVY)

    tb = _textbox(slide, Inches(1.0), Inches(2.4), Inches(11.3), Inches(1.3))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "Exact-Objective Beam Routing"
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = "Calibri"

    tb2 = _textbox(slide, Inches(1.0), Inches(3.55), Inches(11.3), Inches(0.6))
    p2 = tb2.text_frame.paragraphs[0]
    run2 = p2.add_run()
    run2.text = "QSITE 2026 \u2014 Computational Track"
    run2.font.size = Pt(24)
    run2.font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)
    run2.font.name = "Calibri"

    tb3 = _textbox(slide, Inches(1.0), Inches(4.15), Inches(11.3), Inches(0.5))
    p3 = tb3.text_frame.paragraphs[0]
    run3 = p3.add_run()
    run3.text = "Routing and compiling quantum circuits onto a sparse hardware graph"
    run3.font.size = Pt(16)
    run3.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
    run3.font.name = "Calibri"

    tb4 = _textbox(slide, Inches(1.0), Inches(6.6), Inches(11.3), Inches(0.5))
    p4 = tb4.text_frame.paragraphs[0]
    run4 = p4.add_run()
    run4.text = "Official-scorer verified \u00b7 4 of 6 benchmarks proven optimal \u00b7 67.5 vs. 283.5 baseline"
    run4.font.size = Pt(14)
    run4.font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)
    run4.font.italic = True
    run4.font.name = "Calibri"

    _notes(
        slide,
        "[Not part of the timed script -- a few seconds to say who you are and what "
        "this is, then move to the next slide.]\n\n"
        "Hi -- this is our submission for the QSITE 2026 Computational Track: "
        "routing and compiling quantum circuits onto hardware that isn't fully "
        "connected. Here's what we built, and how well it scores.",
    )


def add_content_slide(prs, kicker, title, bullets=None, image=None, notes="",
                       accent=NAVY, bullet_font=20):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, WHITE)
    _kicker_and_title(slide, kicker, title, accent)

    if bullets:
        _bullets(slide, bullets, font_size=bullet_font)
    if image:
        top = Inches(1.95)
        left = Inches(0.5)
        max_w = Inches(12.33)
        max_h = Inches(5.15) if not bullets else Inches(4.35)
        if bullets:
            top = Inches(2.85)
        _image_fitted(slide, ASSETS / image, top, left, max_w, max_h)

    _notes(slide, notes)
    return slide


SLIDES = [
    # -- 0:00-0:35 -----------------------------------------------------
    dict(
        kicker="0:00 \u2013 0:35",
        title="The Problem",
        bullets=[
            "20 qubits on the chip, each wired to only 1, 2, or 3 neighbors",
            "A two-qubit gate can only run on a wire that actually exists",
            "No wire between them? Insert SWAPs to walk the qubits together",
            ("Score: 1 SWAP = 1 point \u00b7 1 parallel time step = 0.5 points", AMBER),
            "6 programs, one combined total \u2014 lower is better",
        ],
        notes=(
            "You are compiling a quantum program onto a chip that is not fully "
            "connected. Twenty qubits. Each one is wired to one, two, or three "
            "neighbors. A two-qubit gate can only run across a wire that "
            "actually exists.\n\n"
            "If the two qubits are not neighbors, you insert SWAPs and walk "
            "them together. Every SWAP costs one point. Every parallel time "
            "step costs half a point. Six programs, one total. Lower wins."
        ),
    ),
    # -- 0:35-1:45 -------------------------------------------------------
    dict(
        kicker="0:35 \u2013 1:45",
        title="What We Did \u2014 Three Ideas",
        bullets=[
            ("1. Zero-SWAP embedding \u2014 some circuits already fit the chip's shape", GREEN),
            "2. Beam search on the real score \u2014 track the exact cost as we go, keep the best few paths",
            "3. Exhaustive search on small cases \u2014 proves some answers are the best possible",
        ],
        bullet_font=22,
        notes=(
            "Three ideas.\n\n"
            "First: sometimes you need zero SWAPs. This chip has a path through "
            "all twenty qubits, and chain and VQE circuits sit on it. Two "
            "benchmarks finish with no SWAPs at all. The score is just half the "
            "circuit's own critical-path depth, and that is a floor on any "
            "hardware.\n\n"
            "Second: when that embedding does not exist, we beam-search the "
            "exact score. The qubits meet along a short path, and each partial "
            "route tracks when every qubit is free, so the cost so far is the "
            "real score. We keep the best few and look ahead. Placement is "
            "lazy: a qubit gets a home the first time it is used. Locking the "
            "placement up front was worse on the dense programs.\n\n"
            "Third: on the smaller instances, a separate exhaustive search "
            "scores routings by that same objective. When our score hits the "
            "bound, the result is optimal."
        ),
    ),
    # -- 1:45-4:10, part 1: scoreboard ------------------------------------
    dict(
        kicker="1:45 \u2013 4:10",
        title="The Results",
        image="scoreboard.png",
        notes=(
            "[Put up scoreboard.png and leave it up for this slide.]\n\n"
            "This table is the official scorer. Twenty-second budget, default "
            "seed. Baseline 283.5, ours 67.5, a 76 percent reduction, from the "
            "same function on every benchmark.\n\n"
            "Four of six benchmarks, including the trickiest small one, are "
            "mathematically proven optimal, not just best-found."
        ),
    ),
    # -- ghz_star ----------------------------------------------------------
    dict(
        kicker="1:45 \u2013 4:10",
        title="ghz_star \u2014 Proven Optimal",
        image="before_after_ghz_star.png",
        accent=GREEN,
        notes=(
            "[Switch to before_after_ghz_star.png.]\n\n"
            "ghz_star. Baseline 14: seven SWAPs, depth 14. Ours, 6.5: two "
            "SWAPs, depth 9. Red edges are the SWAPs. The hub has seven leaves "
            "and the chip's max degree is three, so it cannot sit next to "
            "everyone. Two SWAPs is what that argument requires, and the score "
            "lands on 6.5 exactly. That is the proof.\n\n"
            "[Optional, only if you have 10 spare seconds: play "
            "routing_ghz_star.gif here -- one frame per second, initial "
            "placement then each SWAP and gate. Skip it if you're behind on "
            "time.]"
        ),
    ),
    # -- chain / vqe / ladder ------------------------------------------------
    dict(
        kicker="1:45 \u2013 4:10",
        title="Three More, Proven Optimal",
        bullets=[
            ("chain_trotter \u2014 4.5, zero SWAPs, proven", GREEN),
            ("vqe_layers \u2014 3.0, zero SWAPs, proven (baseline was 58)", GREEN),
            ("ladder_trotter \u2014 6.5: 3 SWAPs, depth 7, proven by exhaustive search", GREEN),
        ],
        accent=GREEN,
        bullet_font=22,
        notes=(
            "chain_trotter is 4.5, zero SWAPs, proven. vqe_layers is 3.0, zero "
            "SWAPs, proven. The baseline on vqe was 58.\n\n"
            "ladder_trotter is 6.5: three SWAPs, depth 7. An exhaustive search, "
            "separate from the solver, proved no valid routing scores below "
            "6.5. Our solution matches that bound, and the official scorer "
            "validates it."
        ),
    ),
    # -- qaoa_random ---------------------------------------------------------
    dict(
        kicker="1:45 \u2013 4:10",
        title="qaoa_random \u2014 Best Found",
        bullets=[
            ("Best found: 11.5 \u2014 6 SWAPs, depth 11 (baseline was 39)", AMBER),
            "Proven floor: 9.0 \u2014 at least 5 SWAPs required, depth already at least 8",
            ("At most 2.5 points still open", AMBER),
        ],
        accent=AMBER,
        bullet_font=22,
        notes=(
            "qaoa_random is best found at 11.5, against a baseline of 39 and a "
            "proven floor of 9. At least five SWAPs are required and the depth "
            "is already at least 8, so at most two and a half points remain "
            "open."
        ),
    ),
    # -- dense_random ---------------------------------------------------------
    dict(
        kicker="1:45 \u2013 4:10",
        title="dense_random \u2014 Still Open",
        image="before_after_dense_random.png",
        accent=AMBER,
        notes=(
            "[Switch to before_after_dense_random.png.]\n\n"
            "dense_random is still open. Forty interactions on fourteen "
            "qubits. The baseline inserts 90 SWAPs at depth 64, score 122. We "
            "got 24 SWAPs, depth 23, score 35.5. The proven floor is 17: at "
            "least eleven SWAPs, and depth at least 12. That floor comes from "
            "a published SAT encoding of a more permissive routing model, so "
            "it is a real floor under our stricter rules. Our current search "
            "plateaus at 35.5. A 240-trial sweep, a 650-run randomized-lineage "
            "sweep, and about three thousand officially scored window splices "
            "all stop there. Whether the true minimum is close to 17 or close "
            "to 35.5 is still an open question."
        ),
    ),
    # -- 4:10-4:35: close ------------------------------------------------------
    dict(
        kicker="4:10 \u2013 4:35",
        title="Where We Landed",
        image="scoreboard.png",
        notes=(
            "[Back to scoreboard.png.]\n\n"
            "So: 67.5 versus 283.5. Four benchmarks match a proof. qaoa is "
            "inside two and a half points of its floor. dense is still open "
            "between 17 and 35.5. Every answer is checked by the official "
            "scorer before we accept it."
        ),
    ),
    # -- stretch bonus ------------------------------------------------------
    dict(
        kicker="4:10 \u2013 4:35",
        title="Bonus: Fewer Gates, More Points",
        bullets=[
            "Rewrote every circuit into the native gate set (RZ, SX, CNOT)",
            ("410 fewer gates than the provided baseline decomposer (650 \u2192 240)", GREEN),
            ("Organizers confirmed this bonus counts \u2014 that's 41.0 points", GREEN),
            ("Reported total: 67.5 \u2212 41.0 = 26.5", GREEN),
        ],
        accent=GREEN,
        bullet_font=22,
        notes=(
            "One more number: the organizers confirmed our stretch-goal gate "
            "count counts too. Rewritten into native gates, our circuits use "
            "410 fewer gates than the provided bad decomposer. That's "
            "forty-one more points, so the total we're reporting is 26.5."
        ),
    ),
]

APPENDIX_QA = [
    (
        "Is dense_random optimal?",
        "No. Proven floor 17.0, our score 35.5: 24 SWAPs, depth 23, baseline "
        "122. The floor is at least eleven SWAPs, from a published SAT "
        "encoding of a more permissive model, plus depth at least 12. It "
        "rests on that encoding. A 240-trial sweep, 650 randomized lineages, "
        "and about three thousand window splices all plateau at 35.5, so "
        "this search approach has stalled. That does not locate the true "
        "minimum. It may sit near 17, or near 35.5. We have not settled "
        "which.",
    ),
    (
        "Is qaoa_random optimal?",
        "No. Best found is 11.5: 6 SWAPs, depth 11, baseline 39. The proven "
        "floor is 9.0, because at least five SWAPs are required and the "
        "critical-path depth is 8. Headroom is at most 2.5. A search aimed "
        "at ruling out 11.0 and 10.5 timed out before it finished, so "
        "\"within half a point\" is not a claim we can make.",
    ),
    (
        "Will the grader get 67.5?",
        "With the default seed and the 20-second budget, yes, that is what "
        "the official scorer returned on the run checked for this demo. It "
        "is an anytime search. Under a heavy machine it can stop on an "
        "earlier checkpoint. We saw qaoa_random come back 12.5 instead of "
        "11.5 once, before we raised the budget from 10 seconds to 20. A "
        "different seed can land a point or two worse on the hard "
        "benchmarks. It will not come back invalid, and it will not go "
        "below the proven floors.",
    ),
    (
        "Did you do the stretch goals \u2014 decomposition and single-qubit "
        "optimization?",
        "Yes. decompose.py rewrites SWAP into 3 CNOTs and each program "
        "two-qubit gate into 1 CNOT, into the native RZ, SX, CNOT set; "
        "optimize_1q.py merges and cancels the single-qubit runs that "
        "produces. Checked against the deleted reference decomposer, which "
        "pads every op with RZ(0) identities -- seven gates per SWAP, three "
        "per two-qubit gate -- our circuits use 240 gates instead of 650. "
        "That's 410 gates saved, 41.0 points at the README's N times 0.1. "
        "verify_stretch.py reproduces that count and checks the "
        "decomposition against the routed program structurally.\n\n"
        "The README still lists decompose() and optimize_1q() as optional "
        "and still says the score improves by N times 0.1, even though an "
        "earlier commit removed that bonus from the headline formula and "
        "deleted the reference implementation. We asked the organizers "
        "directly: it counts. That is where the 26.5 comes from.",
    ),
    (
        "How do you know ladder_trotter is fully optimal?",
        "An exhaustive search over placements and SWAP sequences, with the "
        "real swaps-plus-half-depth score, proved that no valid routing "
        "scores below 6.5. Our 3-SWAP, depth-7 solution matches that bound, "
        "and the official scorer validates it. The same search code was "
        "cross-checked against brute force on 65 small instances with zero "
        "disagreements. The old caveat, \"three SWAPs is minimal but a "
        "shallower depth might still win,\" is closed.",
    ),
]


def add_thank_you_slide(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, NAVY)

    tb = _textbox(slide, Inches(1.0), Inches(2.9), Inches(11.3), Inches(1.0))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Thank You"
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = "Calibri"

    tb2 = _textbox(slide, Inches(1.0), Inches(3.95), Inches(11.3), Inches(0.6))
    p2 = tb2.text_frame.paragraphs[0]
    run2 = p2.add_run()
    run2.text = "Happy to take questions."
    run2.font.size = Pt(20)
    run2.font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)
    run2.font.name = "Calibri"

    _notes(slide, "Happy to take questions.")


def add_appendix_divider(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, PANEL)
    tb = _textbox(slide, Inches(1.0), Inches(2.9), Inches(11.3), Inches(1.0))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Appendix \u2014 Anticipated Questions"
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.color.rgb = INK
    run.font.name = "Calibri"

    tb2 = _textbox(slide, Inches(1.0), Inches(3.75), Inches(11.3), Inches(0.9))
    p2 = tb2.text_frame.paragraphs[0]
    run2 = p2.add_run()
    run2.text = (
        "Not part of the timed 5-minute script. Use these only if you want "
        "to pre-empt a question in your recording -- each note below is "
        "written to be read exactly as-is."
    )
    run2.font.size = Pt(16)
    run2.font.italic = True
    run2.font.color.rgb = GRAY
    run2.font.name = "Calibri"

    _notes(
        slide,
        "[Divider slide -- nothing to read aloud. The next few slides are "
        "optional backup material, only needed if you choose to address "
        "likely questions directly in the recording.]",
    )


def add_qa_slide(prs, question: str, answer: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, WHITE)

    # Questions vary a lot in length, so give the title generous, fixed
    # vertical room for up to two wrapped lines (rather than guessing the
    # wrap point) and shrink the font on long questions so wrapping is rare.
    kb = _textbox(slide, Inches(0.7), Inches(0.35), Inches(8), Inches(0.4))
    krun = kb.text_frame.paragraphs[0].add_run()
    krun.text = "Q&A (OPTIONAL)"
    krun.font.size = Pt(14)
    krun.font.bold = True
    krun.font.color.rgb = GRAY
    krun.font.name = "Calibri"

    title_size = 30 if len(question) <= 55 else 24
    tb = _textbox(slide, Inches(0.7), Inches(0.72), Inches(11.9), Inches(1.5))
    trun = tb.text_frame.paragraphs[0].add_run()
    trun.text = question
    trun.font.size = Pt(title_size)
    trun.font.bold = True
    trun.font.color.rgb = INK
    trun.font.name = "Calibri"

    line = slide.shapes.add_shape(1, Inches(0.7), Inches(2.05), Inches(2.0), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = GRAY
    line.line.fill.background()
    line.shadow.inherit = False

    box = _textbox(slide, Inches(0.85), Inches(2.3), Inches(11.6), Inches(4.4))
    tf = box.text_frame
    tf.word_wrap = True
    for i, para in enumerate(answer.split("\n\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.3
        p.space_after = Pt(12)
        run = p.add_run()
        run.text = para
        run.font.size = Pt(18)
        run.font.color.rgb = INK
        run.font.name = "Calibri"

    _notes(slide, answer)


def main() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_title_slide(prs)
    for spec in SLIDES:
        add_content_slide(
            prs,
            kicker=spec["kicker"],
            title=spec["title"],
            bullets=spec.get("bullets"),
            image=spec.get("image"),
            notes=spec["notes"],
            accent=spec.get("accent", NAVY),
            bullet_font=spec.get("bullet_font", 20),
        )
    add_thank_you_slide(prs)
    add_appendix_divider(prs)
    for question, answer in APPENDIX_QA:
        add_qa_slide(prs, question, answer)

    prs.save(OUT)
    print(f"wrote {OUT}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
